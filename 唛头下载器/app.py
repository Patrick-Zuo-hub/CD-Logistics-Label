# -*- coding: utf-8 -*-
"""物流单据批量下载器
按物流单号批量下载：
  · 唛头 PDF  （getOutMarkPdf）
  · POD 签收附件（getlocusinfo 轨迹里的 pics 附件）
输入方式：上传 Excel 或 文本框粘贴单号。

凭证读取优先级：
  1. 环境变量 CLIENT_ID / CLIENT_SECRET（公网部署，凭证不进代码）
  2. 同目录 env.doc 文件（本地运行）

无状态设计：结果 zip 直接随响应返回，不写服务器磁盘，可直接部署 Vercel。
本地运行时额外把文件存到「唛头PDF」/「POD」文件夹。
"""
import base64
import io
import os
import re
import socket
import zipfile
from pathlib import Path

import requests
from flask import Flask, request, render_template_string, jsonify
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
API_HOST = "http://api.ywcd56.com"
MARK_URL = f"{API_HOST}/openapi/getOutMarkPdf"
LOCUS_URL = f"{API_HOST}/openapi/getlocusinfo"
MARK_DIR = BASE_DIR / "唛头PDF"
POD_DIR = BASE_DIR / "POD"
IS_SERVERLESS = bool(os.environ.get("VERCEL") or os.environ.get("SERVERLESS"))

app = Flask(__name__)


def lan_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except OSError:
        return "127.0.0.1"


def load_env():
    cid = os.environ.get("CLIENT_ID", "").strip()
    sec = os.environ.get("CLIENT_SECRET", "").strip()
    if cid and sec:
        return cid, sec
    creds = {}
    env_path = BASE_DIR / "env.doc"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            creds[k.strip().upper()] = v.strip()
    return cid or creds.get("CLIENT_ID", ""), sec or creds.get("CLIENT_SECRET", "")


WAYBILL_RE = re.compile(r"^[A-Za-z0-9\-_]+$")


def _dedup(nums):
    seen, out = set(), []
    for n in nums:
        n = n.strip()
        if n and WAYBILL_RE.match(n) and n not in seen:
            seen.add(n)
            out.append(n)
    return out


def waybills_from_excel(file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = [[("" if c is None else str(c).strip()) for c in row]
            for row in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        return []
    col, start = 0, 0
    for i, h in enumerate(rows[0]):
        if h.lower() == "waybill_num" or "单号" in h:
            col, start = i, 1
            break
    return _dedup(row[col] for row in rows[start:] if col < len(row))


def waybills_from_text(text):
    return _dedup(re.split(r"[\s,，;；、]+", text or ""))


def waybills_from_request():
    """从上传的Excel或文本框获取单号列表"""
    file = request.files.get("file")
    if file and file.filename:
        return waybills_from_excel(file.read())
    return waybills_from_text(request.form.get("text", ""))


def _api_get(url, waybill, api_key, token, **kw):
    headers = {"x-api-key": api_key, "token": token}
    return requests.get(url, params={"waybill_num": waybill},
                        headers=headers, timeout=30, **kw)


def fetch_mark(waybill, api_key, token):
    """获取唛头PDF，返回 (pdf_bytes, error)"""
    try:
        r = _api_get(MARK_URL, waybill, api_key, token)
    except requests.RequestException as e:
        return None, f"请求失败: {e}"
    ct = r.headers.get("Content-Type", "")
    if "pdf" in ct or r.content[:4] == b"%PDF":
        return r.content, None
    try:
        data = r.json()
    except ValueError:
        return None, f"响应无法解析 (HTTP {r.status_code})"
    if data.get("code") not in (0, 200, "0", "200"):
        msg = data.get("msg", "")
        return None, f"无唛头数据{f'（{msg}）' if msg else ''}"
    path = data.get("path")
    if not path and isinstance(data.get("data"), dict):
        path = data["data"].get("path")
    if not path:
        return None, "接口未返回PDF地址"
    url = path if path.startswith("http") else API_HOST + ("" if path.startswith("/") else "/") + path
    try:
        pr = requests.get(url, headers={"x-api-key": api_key, "token": token}, timeout=60)
        pr.raise_for_status()
    except requests.RequestException as e:
        return None, f"下载PDF失败: {e}"
    if pr.content[:4] != b"%PDF":
        return None, "下载内容不是PDF"
    return pr.content, None


def _collect_pics(data):
    """从轨迹data里递归收集所有 pics 附件地址"""
    urls = []

    def walk(node):
        if isinstance(node, dict):
            pics = node.get("pics")
            if isinstance(pics, str) and pics.strip():
                for u in re.split(r"[,，;；\s]+", pics.strip()):
                    if u:
                        urls.append(u)
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)

    walk(data)
    # 去重保持顺序
    seen, out = set(), []
    for u in urls:
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def _ext_for(url, content):
    m = re.search(r"\.(pdf|jpe?g|png|gif|webp|bmp|tiff?)(?:\?|$)", url, re.I)
    if m:
        return "." + m.group(1).lower().replace("jpeg", "jpg")
    if content[:4] == b"%PDF":
        return ".pdf"
    if content[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if content[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    return ".jpg"


def fetch_pod(waybill, api_key, token):
    """查轨迹并下载POD附件，返回 (list[(filename, bytes)], error)"""
    try:
        r = _api_get(LOCUS_URL, waybill, api_key, token)
    except requests.RequestException as e:
        return [], f"请求失败: {e}"
    try:
        data = r.json()
    except ValueError:
        return [], f"响应无法解析 (HTTP {r.status_code})"
    if data.get("code") not in (0, 200, "0", "200"):
        msg = data.get("msg", "")
        return [], f"查询失败{f'（{msg}）' if msg else ''}"
    pics = _collect_pics(data.get("data", []))
    if not pics:
        return [], "无POD附件"
    files, errs = [], 0
    for i, u in enumerate(pics, 1):
        url = u if u.startswith("http") else API_HOST + ("" if u.startswith("/") else "/") + u
        try:
            pr = requests.get(url, headers={"x-api-key": api_key, "token": token}, timeout=60)
            pr.raise_for_status()
            content = pr.content
        except requests.RequestException:
            errs += 1
            continue
        ext = _ext_for(url, content)
        name = f"{waybill}{ext}" if len(pics) == 1 else f"{waybill}_{i}{ext}"
        files.append((name, content))
    if not files:
        return [], "POD附件下载失败"
    return files, None


PAGE = """<!doctype html><html lang="zh"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>物流单据批量下载器</title>
<style>
body{font-family:-apple-system,"PingFang SC",sans-serif;max-width:660px;margin:36px auto;padding:0 20px;color:#333}
h1{font-size:22px;margin-bottom:4px}
.tabs{display:flex;gap:8px;margin-top:16px}
.tab{flex:1;text-align:center;padding:10px;border:1px solid #ddd;border-radius:10px 10px 0 0;background:#f6f6f6;cursor:pointer;font-size:15px;border-bottom:none}
.tab.active{background:#fff;font-weight:600;color:#2563eb;border-color:#ddd}
.card{border:1px solid #ddd;border-radius:0 10px 10px 10px;padding:24px}
textarea{width:100%;box-sizing:border-box;min-height:90px;border:1px solid #ccc;border-radius:8px;padding:10px;font-size:14px;font-family:inherit}
input[type=file]{margin:6px 0}
.or{color:#999;font-size:13px;margin:10px 0 4px}
button{background:#2563eb;color:#fff;border:0;border-radius:8px;padding:10px 24px;font-size:15px;cursor:pointer;margin-top:12px}
button:disabled{background:#aaa}
.log{white-space:pre-wrap;font-size:13px;background:#f6f6f6;border-radius:8px;padding:12px;margin-top:16px;display:none}
.warn{color:#b45309}.ok{color:#15803d}.err{color:#b91c1c}
.tip{background:#fef9c3;border:1px solid #fde047;border-radius:8px;padding:8px 12px;font-size:13px;color:#854d0e;margin:0 0 12px}
a.dl{display:inline-block;margin-top:12px;color:#2563eb;font-weight:600}
</style></head><body>
<h1>物流单据批量下载器</h1>
{% if not has_creds %}<p class="tip">⚠️ 尚未配置凭证：请设置环境变量 CLIENT_ID / CLIENT_SECRET（部署平台）或填写 env.doc（本地）。</p>{% endif %}
<div class="tabs">
  <div class="tab active" data-t="mark" onclick="sw('mark')">下载唛头</div>
  <div class="tab" data-t="pod" onclick="sw('pod')">下载 POD</div>
</div>
<div class="card">
  <p class="tip">💡 建议每次查询单量小于 <b>60 单</b>；单量较大时请自行分批，以免请求超时。</p>
  <form class="job" data-kind="mark">
    <p>方式一：上传含物流单号的 Excel（.xlsx）</p>
    <input type="file" name="file" accept=".xlsx,.xlsm">
    <p class="or">方式二：直接粘贴单号（换行、逗号或空格分隔均可）</p>
    <textarea name="text" placeholder="CD2607200276&#10;CD2607200277&#10;CD2607200279"></textarea>
    <br><button type="submit">开始下载唛头</button>
    <div class="log"></div>
  </form>
  <form class="job" data-kind="pod" style="display:none">
    <p>方式一：上传含物流单号的 Excel（.xlsx）</p>
    <input type="file" name="file" accept=".xlsx,.xlsm">
    <p class="or">方式二：直接粘贴单号（换行、逗号或空格分隔均可）</p>
    <textarea name="text" placeholder="CD2607200276&#10;CD2607200277"></textarea>
    <br><button type="submit">开始下载 POD</button>
    <div class="log"></div>
  </form>
</div>
<script>
function sw(t){
  document.querySelectorAll('.tab').forEach(e=>e.classList.toggle('active',e.dataset.t===t));
  document.querySelectorAll('.job').forEach(f=>f.style.display=f.dataset.kind===t?'block':'none');
}
document.querySelectorAll('.job').forEach(form=>{
  const kind=form.dataset.kind, log=form.querySelector('.log'), btn=form.querySelector('button');
  form.addEventListener('submit',async e=>{
    e.preventDefault();
    const hasFile=form.file.files.length>0, hasText=form.text.value.trim().length>0;
    if(!hasFile&&!hasText){alert('请上传 Excel 或在文本框粘贴单号');return;}
    btn.disabled=true;log.style.display='block';log.textContent='处理中，请稍候…';
    try{
      const r=await fetch('/process/'+kind,{method:'POST',body:new FormData(form)});
      const d=await r.json();
      log.innerHTML=d.log.map(l=>{
        const cls=l.startsWith('✔')?'ok':(l.startsWith('✘')?'err':(l.startsWith('⚠')?'warn':''));
        return cls?`<span class="${cls}">${l}</span>`:l;
      }).join('\\n');
      if(d.zip_b64){
        const bin=atob(d.zip_b64),arr=new Uint8Array(bin.length);
        for(let i=0;i<bin.length;i++)arr[i]=bin.charCodeAt(i);
        const url=URL.createObjectURL(new Blob([arr],{type:'application/zip'}));
        const a=document.createElement('a');a.href=url;a.className='dl';
        a.download=(kind==='pod'?'POD':'唛头PDF')+'.zip';a.textContent='📦 下载全部 (zip)';
        log.appendChild(document.createElement('br'));log.appendChild(a);
      }
    }catch(err){log.textContent='出错: '+err;}
    btn.disabled=false;
  });
});
</script></body></html>"""


@app.route("/")
def index():
    api_key, secret = load_env()
    return render_template_string(PAGE, has_creds=bool(api_key and secret))


def _run(kind):
    api_key, secret = load_env()
    log = []
    if not (api_key and secret):
        return jsonify(log=["✘ 未配置凭证：请设置环境变量 CLIENT_ID / CLIENT_SECRET（或本地 env.doc）"], zip_b64=None)
    try:
        waybills = waybills_from_request()
    except Exception as e:
        return jsonify(log=[f"✘ 单号解析失败: {e}"], zip_b64=None)
    if not waybills:
        return jsonify(log=["✘ 未获取到任何物流单号"], zip_b64=None)
    log.append(f"共读取 {len(waybills)} 个单号")
    if len(waybills) > 60:
        log.append(f"⚠️ 本次 {len(waybills)} 单已超过建议上限（60单），若超时请分批（每批≤60单）重试")

    out_dir = MARK_DIR if kind == "mark" else POD_DIR
    if not IS_SERVERLESS:
        out_dir.mkdir(exist_ok=True)
    ok, results, failed = 0, [], []

    for wb in waybills:
        if kind == "mark":
            pdf, err = fetch_mark(wb, api_key, secret)
            items = [(f"{wb}.pdf", pdf)] if pdf else []
        else:
            items, err = fetch_pod(wb, api_key, secret)
        if items:
            for name, data in items:
                if not IS_SERVERLESS:
                    (out_dir / name).write_bytes(data)
                results.append((name, data))
            cnt = f"（{len(items)}个附件）" if kind == "pod" and len(items) > 1 else ""
            log.append(f"✔ {wb} 已下载{cnt}")
            ok += 1
        else:
            failed.append(wb)
            log.append(f"✘ {wb} {err}")

    log.append(f"\n完成：成功 {ok} / {len(waybills)}")
    if failed:
        log.append(f"✘ 失败单号（{len(failed)}个）: {'、'.join(failed)}")
    zip_b64 = None
    if results:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, data in results:
                zf.writestr(name, data)
        zip_b64 = base64.b64encode(buf.getvalue()).decode()
    return jsonify(log=log, zip_b64=zip_b64)


@app.route("/process/mark", methods=["POST"])
def process_mark():
    return _run("mark")


@app.route("/process/pod", methods=["POST"])
def process_pod():
    return _run("pod")


# 兼容旧路径
@app.route("/process", methods=["POST"])
def process_legacy():
    return _run("mark")


if __name__ == "__main__":
    ip = lan_ip()
    print("物流单据批量下载器已启动")
    print(f"  本机访问:   http://127.0.0.1:8756")
    print(f"  局域网访问: http://{ip}:8756  （发给同一网络的同事即可）")
    app.run(host="0.0.0.0", port=8756, threaded=True)
